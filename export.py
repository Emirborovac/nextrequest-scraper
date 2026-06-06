import sys, store
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


def run(out="oakland_crashes.xlsx"):
    c = store.connect()
    rows = c.execute("SELECT * FROM docs WHERE is_crash='yes' ORDER BY created_at DESC").fetchall()
    c.close()
    wb = Workbook()
    ws = wb.active
    ws.title = "Crash reports"
    headers = ["Document"] + [store.EXCEL_LABELS[f] for f in store.EXCEL_FIELDS] + ["Uploaded", "Source Link"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0B3D91")
    for r in rows:
        ws.append([r["title_file"]] + [r[f] for f in store.EXCEL_FIELDS] + [r["created_at"], r["download_url"]])
    wb.save(out)
    print("wrote", out, "| crash rows:", len(rows))


if __name__ == "__main__":
    run(sys.argv[1] if len(sys.argv) > 1 else "oakland_crashes.xlsx")
