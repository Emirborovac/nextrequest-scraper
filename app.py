import io, os, csv, time, sqlite3
try:
    from dotenv import load_dotenv
    load_dotenv()                      # read config from a .env file if present
except Exception:
    pass
from flask import Flask, render_template_string, send_file, abort, request
import store, nextrequest as nr, ma_csv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

app = Flask(__name__)
store.init()

PAGE = """<!doctype html>
<html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">
<title>NextRequest Crash Monitor</title>
{% if refresh %}<meta http-equiv="refresh" content="{{ refresh }}">{% endif %}
<style>
:root{--navy:#0b3d91;--ink:#1c2b3a;--muted:#6b7c91;--line:#dfe7f0;--red:#b23b3b;--pale:#fdeaea;}
*{box-sizing:border-box}body{margin:0;font-family:-apple-system,"Segoe UI",Roboto,Arial,sans-serif;color:var(--ink);background:#fff;font-size:15px}
.page{max-width:1320px;margin:0 auto;padding:28px 36px 60px}
h1{color:var(--navy);font-size:1.5rem;margin:0 0 2px}h2{color:var(--navy);font-size:1.1rem;margin:26px 0 8px}
.sub{color:var(--muted);margin:0 0 18px;font-size:.9rem}
.cards{display:flex;gap:14px;flex-wrap:wrap;margin:10px 0 6px}
.card{border:1px solid var(--line);border-radius:8px;padding:14px 20px;min-width:130px}
.card .n{font-size:1.8rem;font-weight:700;color:var(--navy)}.card .l{color:var(--muted);font-size:.82rem}
.card.crash{border-color:var(--red)}.card.crash .n{color:var(--red)}
.meta{color:var(--muted);font-size:.82rem;margin:8px 0 14px}
a.btn{display:inline-block;background:var(--navy);color:#fff;text-decoration:none;padding:9px 16px;border-radius:6px;font-weight:600;font-size:.9rem}
a.btn:hover{background:#0a2f73}
table{width:100%;border-collapse:collapse;font-size:.84rem;margin-top:6px}
th{text-align:left;color:var(--navy);border-bottom:2px solid var(--navy);padding:7px 9px;white-space:nowrap;position:sticky;top:0;background:#fff}
td{padding:6px 9px;border-bottom:1px solid var(--line);vertical-align:top;overflow-wrap:anywhere}
tr.yes{background:var(--pale)}
.wrap{overflow-x:auto;border:1px solid var(--line);border-radius:8px;margin-top:6px}
.empty{color:var(--muted);text-align:center;padding:18px}
a{color:var(--navy)}
.nav{display:flex;gap:6px;margin:4px 0 18px;border-bottom:2px solid var(--line)}
.nav a{padding:8px 16px;text-decoration:none;color:var(--muted);font-weight:600;border-bottom:3px solid transparent;margin-bottom:-2px}
.nav a.on{color:var(--navy);border-bottom-color:var(--navy)}
</style></head><body><div class="page">
<h1>Crash Data Console</h1>
<div class="nav"><a href="/" class="on">Oakland &middot; NextRequest</a><a href="/massachusetts">Massachusetts</a></div>
<p class="sub">Source: {{ base }} &middot; monitoring stream: &ldquo;{{ term or 'all' }}&rdquo;{% if refresh %} &middot; auto-refreshes every {{ refresh }}s{% endif %}</p>
<div class="cards">
  <div class="card"><div class="n">{{ "{:,}".format(stats.total) }}</div><div class="l">docs scanned</div></div>
  <div class="card crash"><div class="n">{{ "{:,}".format(stats.crash) }}</div><div class="l">crash reports found</div></div>
  <div class="card"><div class="n">{{ "{:,}".format(stats.done) }}</div><div class="l">classified</div></div>
  <div class="card"><div class="n">{{ "{:,}".format(stats.pending) }}</div><div class="l">pending</div></div>
</div>
<p class="meta">Upload date range: <b>{{ dr[0] or '—' }} &rarr; {{ dr[1] or '—' }}</b>
  &middot; statuses: {% for k,v in by_status.items() %}{{ k }}={{ v }}{% if not loop.last %}, {% endif %}{% endfor %}</p>
<a class="btn" href="/export.xlsx">&#8595; Export crash reports (Excel)</a>

<h2>Scraping operations</h2>
<div class="wrap"><table><thead><tr><th>#</th><th>Type</th><th>Started (UTC)</th><th>Finished (UTC)</th>
<th>New docs</th><th>Parsed</th><th>Crashes</th><th>Status</th></tr></thead><tbody>
{% for r in runs %}<tr class="{{ 'yes' if r.crashes and r.crashes>0 else '' }}">
<td>{{ r.id }}</td><td>{{ r.kind }}</td><td>{{ r.started_at }}</td>
<td>{{ r.finished_at if r.finished_at else 'running…' }}</td>
<td>{{ '' if r.new_docs is none else r.new_docs }}</td>
<td>{{ '' if r.classified is none else r.classified }}</td>
<td>{{ '' if r.crashes is none else r.crashes }}</td><td>{{ r.status }}</td></tr>{% endfor %}
{% if not runs %}<tr><td colspan="8" class="empty">No runs yet — start one with <code>run.py seed</code> or <code>run.py monitor</code>.</td></tr>{% endif %}
</tbody></table></div>

<h2>Crash reports ({{ crashes|length }})</h2>
<div class="wrap"><table><thead><tr><th>Uploaded</th><th>Document</th>
{% for f in fields %}<th>{{ labels[f] }}</th>{% endfor %}<th>Source</th></tr></thead><tbody>
{% for r in crashes %}<tr><td>{{ r.created_at }}</td><td>{{ r.title_file }}</td>
{% for f in fields %}<td>{{ r[f] or '' }}</td>{% endfor %}
<td><a href="{{ r.download_url }}" target="_blank" rel="noopener">PDF</a></td></tr>{% endfor %}
{% if not crashes %}<tr><td colspan="20" class="empty">No crash reports identified yet — the monitor will list them here as they are found.</td></tr>{% endif %}
</tbody></table></div>

<h2>Recent classifications &mdash; why each was / wasn't flagged</h2>
<div class="wrap"><table><thead><tr><th>Uploaded</th><th>Document</th><th>Crash?</th><th>Reason (why / why not a crash report)</th></tr></thead><tbody>
{% for r in recent %}<tr class="{{ 'yes' if r.is_crash=='yes' else '' }}"><td>{{ r.created_at }}</td>
<td>{{ r.title_file }}</td><td>{{ r.is_crash or '' }}</td><td>{{ r.ai_reason or '' }}</td></tr>{% endfor %}
{% if not recent %}<tr><td colspan="4" class="empty">Nothing classified yet — run a seed or wait for the poll.</td></tr>{% endif %}
</tbody></table></div>
</div></body></html>"""


@app.route("/")
def index():
    c = store.connect()
    stats = dict(store.counts())
    by_status = {r["status"]: r["n"] for r in
                 c.execute("SELECT status, COUNT(*) n FROM docs GROUP BY status").fetchall()}
    dr = c.execute("SELECT MIN(created_at), MAX(created_at) FROM docs").fetchone()
    crashes = c.execute("SELECT * FROM docs WHERE is_crash='yes' ORDER BY created_at DESC LIMIT 1000").fetchall()
    recent = c.execute("SELECT created_at,title_file,is_crash,ai_reason FROM docs WHERE status='done' "
                       "ORDER BY classified_at DESC LIMIT 100").fetchall()
    runs = c.execute("SELECT * FROM runs ORDER BY id DESC LIMIT 25").fetchall()
    c.close()
    return render_template_string(PAGE, stats=stats, by_status=by_status, dr=dr,
                                  crashes=crashes, recent=recent, runs=runs,
                                  fields=store.EXCEL_FIELDS, labels=store.EXCEL_LABELS,
                                  base=nr.BASE, term=nr.SEARCH_TERM,
                                  refresh=int(os.environ.get("REFRESH_SECS", "20")))


@app.route("/export.xlsx")
def export():
    c = store.connect()
    rows = c.execute("SELECT * FROM docs WHERE is_crash='yes' ORDER BY created_at DESC").fetchall()
    c.close()
    wb = Workbook(); ws = wb.active; ws.title = "Crash reports"
    headers = ["Document"] + [store.EXCEL_LABELS[f] for f in store.EXCEL_FIELDS] + ["Uploaded", "Source Link"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="0B3D91")
    for r in rows:
        ws.append([r["title_file"]] + [r[f] for f in store.EXCEL_FIELDS] + [r["created_at"], r["download_url"]])
    bio = io.BytesIO(); wb.save(bio); bio.seek(0)
    return send_file(bio, as_attachment=True, download_name="crash_reports.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


MA_DB = os.environ.get("MA_DB", os.path.join(os.path.dirname(os.path.abspath(__file__)), "ma.db"))


def ma_meta():
    if not os.path.exists(MA_DB):
        return None
    c = sqlite3.connect(MA_DB)
    try:
        row = c.execute("SELECT MIN(crash_date_iso), MAX(crash_date_iso), "
                        "COUNT(*), COUNT(DISTINCT CRASH_NUMB) FROM ma").fetchone()
    except Exception:
        c.close(); return None
    c.close()
    if not row or row[0] is None:
        return None
    return {"dmin": row[0], "dmax": row[1], "vehicles": row[2], "accidents": row[3],
            "updated": time.strftime("%Y-%m-%d %H:%M UTC", time.gmtime(os.path.getmtime(MA_DB)))}


@app.route("/ma.csv")
def ma_download():
    meta = ma_meta()
    if not meta:
        abort(404)
    frm = request.args.get("from") or meta["dmin"]
    to = request.args.get("to") or meta["dmax"]
    cols = ma_csv.COLUMNS
    c = sqlite3.connect(MA_DB)
    rows = c.execute('SELECT %s FROM ma WHERE crash_date_iso BETWEEN ? AND ? ORDER BY crash_date_iso DESC'
                     % ",".join('"%s"' % x for x in cols), (frm, to)).fetchall()
    c.close()
    sio = io.StringIO()
    w = csv.writer(sio); w.writerow(cols); w.writerows(rows)
    data = io.BytesIO(sio.getvalue().encode("utf-8")); data.seek(0)
    return send_file(data, as_attachment=True,
                     download_name="massachusetts_%s_to_%s.csv" % (frm, to), mimetype="text/csv")


MA_PAGE = """<!doctype html>
<html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">
<title>Massachusetts — Crash+Vehicle</title>
<style>
:root{--navy:#0b3d91;--ink:#1c2b3a;--muted:#6b7c91;--line:#dfe7f0;}
*{box-sizing:border-box}body{margin:0;font-family:-apple-system,"Segoe UI",Roboto,Arial,sans-serif;color:var(--ink);background:#fff;font-size:15px}
.page{max-width:1080px;margin:0 auto;padding:28px 36px 60px}
h1{color:var(--navy);font-size:1.5rem;margin:0 0 2px}h2{color:var(--navy);font-size:1.05rem;margin:22px 0 8px}
.sub{color:var(--muted);margin:0 0 10px;font-size:.9rem}
.nav{display:flex;gap:6px;margin:4px 0 18px;border-bottom:2px solid var(--line)}
.nav a{padding:8px 16px;text-decoration:none;color:var(--muted);font-weight:600;border-bottom:3px solid transparent;margin-bottom:-2px}
.nav a.on{color:var(--navy);border-bottom-color:var(--navy)}
a.btn{display:inline-block;background:var(--navy);color:#fff;text-decoration:none;padding:11px 20px;border-radius:6px;font-weight:700;font-size:.95rem}
a.btn:hover{background:#0a2f73}
.filter{display:flex;gap:10px;align-items:flex-end;flex-wrap:wrap;margin:6px 0 16px}
.filter label{font-size:.8rem;color:var(--muted);display:flex;flex-direction:column;gap:3px}
.filter input{border:1px solid var(--line);border-radius:6px;padding:7px 9px;font-size:.9rem}
.filter button{background:var(--navy);color:#fff;border:0;border-radius:6px;padding:8px 16px;font-weight:600;cursor:pointer}
.cards{display:flex;gap:14px;flex-wrap:wrap;margin:8px 0}
.card{border:1px solid var(--line);border-radius:8px;padding:12px 18px;min-width:140px}
.card .n{font-size:1.7rem;font-weight:700;color:var(--navy)}.card .l{color:var(--muted);font-size:.8rem}
.meta{color:var(--muted);font-size:.82rem;margin:6px 0 12px}
table{border-collapse:collapse;font-size:.86rem;width:100%;max-width:580px}
th{text-align:left;color:var(--navy);border-bottom:2px solid var(--navy);padding:6px 10px}
td{padding:5px 10px;border-bottom:1px solid var(--line)}
.bar{display:inline-block;height:11px;background:var(--navy);border-radius:2px;vertical-align:middle}
.empty{color:var(--muted);padding:18px}
</style></head><body><div class="page">
<h1>Crash Data Console</h1>
<div class="nav"><a href="/">Oakland &middot; NextRequest</a><a href="/massachusetts" class="on">Massachusetts</a></div>
<p class="sub">MassDOT IMPACT &middot; Crash + Vehicle (VINs) &middot; 52-column CSV &middot; refreshed daily</p>
{% if meta %}
<form class="filter" method="get" action="/massachusetts">
  <label>From<input type="date" name="from" value="{{ frm }}" min="{{ meta.dmin }}" max="{{ meta.dmax }}"></label>
  <label>To<input type="date" name="to" value="{{ to }}" min="{{ meta.dmin }}" max="{{ meta.dmax }}"></label>
  <button type="submit">Apply</button>
</form>
<div class="cards">
  <div class="card"><div class="n">{{ "{:,}".format(tot_accidents) }}</div><div class="l">accidents in range</div></div>
  <div class="card"><div class="n">{{ "{:,}".format(tot_rows) }}</div><div class="l">vehicle rows (CSV)</div></div>
</div>
<a class="btn" href="/ma.csv?from={{ frm }}&amp;to={{ to }}">&#8595; Download CSV for {{ frm }} &rarr; {{ to }}</a>
<p class="meta">Data available {{ meta.dmin }} &rarr; {{ meta.dmax }} &middot; updated {{ meta.updated }}</p>
<h2>Accidents per day</h2>
<table><thead><tr><th>Date</th><th>Accidents</th><th></th></tr></thead><tbody>
{% for d, n in perday %}<tr><td>{{ d }}</td><td>{{ n }}</td><td><span class="bar" style="width:{{ (n * 240 // maxc) if maxc else 0 }}px"></span></td></tr>{% endfor %}
{% if not perday %}<tr><td colspan="3" class="empty">No accidents in this range.</td></tr>{% endif %}
</tbody></table>
{% else %}
<p class="empty">No data yet &mdash; the Massachusetts cache builds automatically each morning.</p>
{% endif %}
</div></body></html>"""


@app.route("/massachusetts")
def massachusetts():
    meta = ma_meta()
    if not meta:
        return render_template_string(MA_PAGE, meta=None)
    frm = request.args.get("from") or meta["dmin"]
    to = request.args.get("to") or meta["dmax"]
    c = sqlite3.connect(MA_DB)
    perday = c.execute("SELECT crash_date_iso, COUNT(DISTINCT CRASH_NUMB) FROM ma "
                       "WHERE crash_date_iso BETWEEN ? AND ? GROUP BY crash_date_iso "
                       "ORDER BY crash_date_iso DESC", (frm, to)).fetchall()
    tot = c.execute("SELECT COUNT(DISTINCT CRASH_NUMB), COUNT(*) FROM ma "
                    "WHERE crash_date_iso BETWEEN ? AND ?", (frm, to)).fetchone()
    c.close()
    maxc = max([p[1] for p in perday], default=1)
    return render_template_string(MA_PAGE, meta=meta, frm=frm, to=to, perday=perday,
                                  tot_accidents=tot[0], tot_rows=tot[1], maxc=maxc)


if __name__ == "__main__":
    # Single entrypoint: `python app.py` runs the scraper (seed-if-empty + 2h poll)
    # in a background thread AND serves the dashboard. Set RUN_MONITOR=0 for UI-only.
    if os.environ.get("RUN_MONITOR", "1") != "0":
        import threading, run
        threading.Thread(target=run.monitor, daemon=True).start()
        print("[app] monitor thread started (seed if empty, then poll every 2h)")
    app.run(host="0.0.0.0", port=int(os.environ.get("PORT", "8080")), use_reloader=False)
